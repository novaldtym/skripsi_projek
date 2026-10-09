import os, sys, datetime
try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass
import pandas as pd
import openpyxl

BASE_DIR = r"d:\SKRIPSI INFORMATIKA"

print("="*80)
print("MERESET DATABASE FORWARD TESTING KE TITIK 0 TRANSAKSI (CLEAN SLATE)")
print("="*80)

# 1. Reset Evaluasi_Skenario_Trade.csv dan evaluasi_skenario_trade.csv
csv_cols = [
    'Ticket Posisi', 'Waktu Open', 'Waktu Close', 'Simbol', 'Tipe', 'Lot',
    'Harga Entry', 'Harga Exit', 'Pips (P/L)', 'Profit ($ USD)', 'Hasil',
    'Model AI', 'Skenario Entry', 'Zona Entry', 'Tren H4 Entry', 'Tren DXY Entry',
    'Kondisi Makro News', 'Probabilitas Model (%)', 'Analisis Diagnostik (Menang/Kalah)',
    'Evaluasi & Pembelajaran', 'Tanggal Evaluasi', 'Close Candle 75m ($)',
    'Validasi Model 75M', 'Pips Model 75M', 'Durasi', 'Stop Loss (SL)',
    'Take Profit (TP)', 'Keterangan / Alasan Exit'
]
df_empty = pd.DataFrame(columns=csv_cols)

for c_name in ["Evaluasi_Skenario_Trade.csv", "evaluasi_skenario_trade.csv"]:
    p = os.path.join(BASE_DIR, c_name)
    df_empty.to_csv(p, index=False)
    print(f"[OK] {c_name} di-reset bersih (0 transaksi).")

# 2. Reset Laporan_Forward_Testing_Model_Terbaru_SMC.xlsx
excel_m15 = os.path.join(BASE_DIR, "Laporan_Forward_Testing_Model_Terbaru_SMC.xlsx")
if os.path.exists(excel_m15):
    wb = openpyxl.load_workbook(excel_m15)
    
    # Reset Trade Log Pure 100 (v4.2)
    if 'Trade Log Pure 100 (v4.2)' in wb.sheetnames:
        ws_log = wb['Trade Log Pure 100 (v4.2)']
        if ws_log.max_row > 1:
            ws_log.delete_rows(2, ws_log.max_row - 1)
        print("[OK] Sheet 'Trade Log Pure 100 (v4.2)' di-reset (0 baris).")

    # Reset Ringkasan Statistik (v4.2)
    if 'Ringkasan Statistik (v4.2)' in wb.sheetnames:
        ws_stat = wb['Ringkasan Statistik (v4.2)']
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S WIB")
        
        updates = {
            2: ("Nama Model", "LightGBM v5.2 SMC/DXY M15 (65 Fitur: SMC + MTF + EMA 9/26 + DXY SMT POI)"),
            4: ("Ambang Batas Keyakinan (Threshold)", "Versi 5.2 (Pure AI Decision >= 60% Normal, >= 65% Sniper)"),
            7: ("Saldo Terkini", "$500.00 USD"),
            8: ("Total Trade Otomatis Selesai", "0 Trade"),
            9: ("Progres Evaluasi (%)", "0.0%"),
            10: ("Jumlah Trade WIN (Murni)", "0 Trade"),
            11: ("Jumlah Trade BEP (Netral)", "0 Trade"),
            12: ("Jumlah Trade LOSS (Rugi)", "0 Trade"),
            13: ("Win Rate Murni (%) [Win/(Win+Loss)]", "0.00%"),
            14: ("Win Rate Total (%) [Win/Total]", "0.00%"),
            15: ("Total Akumulasi Profit ($ USD)", "$+0.00 USD"),
            16: ("Return on Investment (ROI)", "+0.00%"),
            17: ("Rata-rata Profit per Trade", "$0.00 USD"),
            18: ("Profit Factor", "0.00"),
            19: ("Waktu Pembaruan Terakhir", now_str)
        }
        for row_idx, (metric, val) in updates.items():
            ws_stat.cell(row_idx, 1).value = metric
            ws_stat.cell(row_idx, 2).value = val
        print("[OK] Sheet 'Ringkasan Statistik (v4.2)' di-reset ke nilai awal ($500.00, 0 trade).")

    wb.save(excel_m15)
    print(f"[OK] File {excel_m15} berhasil disimpan.")

# 3. Reset Evaluasi_Skenario_Trade.xlsx
excel_eval = os.path.join(BASE_DIR, "Evaluasi_Skenario_Trade.xlsx")
if os.path.exists(excel_eval):
    try:
        wb_ev = openpyxl.load_workbook(excel_eval)
        for sname in wb_ev.sheetnames:
            ws = wb_ev[sname]
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)
        wb_ev.save(excel_eval)
        print(f"[OK] {excel_eval} di-reset (0 baris).")
    except Exception as e:
        print(f"Warning Evaluasi_Skenario_Trade.xlsx: {e}")

print("RESET DATABASE SELESAI DENGAN SEMPURNA!")
